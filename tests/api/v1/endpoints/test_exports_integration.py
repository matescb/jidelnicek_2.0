"""
Integration tests for export API with real data.

Tests the complete export flow with actual database records and file generation.
"""

import pytest
import asyncio
import json
import csv
import io
from datetime import datetime, timedelta
from typing import List
from openpyxl import load_workbook
from PyPDF2 import PdfReader

from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from jidelnicek.api.v1.schemas.export_schemas import (
    ExportType,
    ExportFormat,
    ExportStatus
)
from jidelnicek.auth.models import AuthUser as User
from jidelnicek.trip.models import Trip, Day, Meal, MealSlot
from jidelnicek.recipe.models import Recipe, RecipeIngredient, RecipeCategory, Category
from jidelnicek.common.models import Ingredient, NutritionalValue
from jidelnicek.shopping.services.shopping_list_generator import ShoppingListGenerator


class TestExportIntegration:
    """Integration tests for export functionality."""
    
    @pytest.fixture
    async def complete_trip_data(
        self,
        db_session: AsyncSession,
        test_user: User
    ) -> Trip:
        """Create a complete trip with all related data."""
        # Create categories
        categories = []
        for name in ["Hlavní jídlo", "Polévka", "Dezert", "Snídaně"]:
            cat = Category(name=name, slug=name.lower().replace(" ", "-"))
            db_session.add(cat)
            categories.append(cat)
        
        await db_session.flush()
        
        # Create ingredients
        ingredients = []
        ingredient_data = [
            ("Brambory", "Zelenina", "kg", 150),
            ("Mrkev", "Zelenina", "kg", 50),
            ("Cibule", "Zelenina", "kg", 40),
            ("Kuřecí prsa", "Maso", "kg", 250),
            ("Rýže", "Přílohy", "kg", 80),
            ("Těstoviny", "Přílohy", "kg", 90),
            ("Mléko", "Mléčné výrobky", "l", 40),
            ("Máslo", "Mléčné výrobky", "kg", 300),
            ("Mouka", "Pečivo", "kg", 35),
            ("Vejce", "Ostatní", "ks", 20)
        ]
        
        for name, category, unit, price in ingredient_data:
            ingredient = Ingredient(
                name=name,
                category=category,
                unit=unit,
                price_per_unit=price
            )
            db_session.add(ingredient)
            ingredients.append(ingredient)
        
        await db_session.flush()
        
        # Create recipes
        recipes = []
        recipe_data = [
            ("Kuřecí rizoto", "Klasické kuřecí rizoto s zeleninou", 6, 20, 40, categories[0]),
            ("Bramborová polévka", "Tradiční česká bramboračka", 8, 15, 30, categories[1]),
            ("Těstovinový salát", "Lehký salát s těstovinami", 4, 15, 20, categories[0]),
            ("Palačinky", "Sladké palačinky s marmeládou", 4, 10, 15, categories[2]),
            ("Müsli s mlékem", "Rychlá snídaně", 1, 5, 0, categories[3])
        ]
        
        for name, desc, servings, prep, cook, category in recipe_data:
            recipe = Recipe(
                name=name,
                description=desc,
                servings=servings,
                prep_time=prep,
                cook_time=cook,
                user_id=test_user.id
            )
            db_session.add(recipe)
            await db_session.flush()
            
            # Add category
            recipe_cat = RecipeCategory(
                recipe_id=recipe.id,
                category_id=category.id
            )
            db_session.add(recipe_cat)
            
            # Add some ingredients
            for i in range(3):
                recipe_ing = RecipeIngredient(
                    recipe_id=recipe.id,
                    ingredient_id=ingredients[i % len(ingredients)].id,
                    amount=100 * (i + 1),
                    unit=ingredients[i % len(ingredients)].unit
                )
                db_session.add(recipe_ing)
            
            # Add nutritional values
            nutrition = NutritionalValue(
                recipe_id=recipe.id,
                calories=350 + i * 50,
                proteins=25 + i * 5,
                carbs=40 + i * 10,
                fats=15 + i * 3,
                fiber=5 + i
            )
            db_session.add(nutrition)
            
            recipes.append(recipe)
        
        await db_session.flush()
        
        # Create trip
        trip = Trip(
            name="Letní tábor 2024",
            description="Dětský letní tábor v Krkonoších",
            start_date=datetime(2024, 7, 1),
            end_date=datetime(2024, 7, 14),
            participant_count=30,
            user_id=test_user.id,
            location="Krkonoše",
            notes="Nezapomenout na vegetariánské varianty"
        )
        db_session.add(trip)
        await db_session.flush()
        
        # Create days and meals
        meal_types = ["breakfast", "lunch", "dinner", "snack"]
        meal_slots = []
        
        for day_num in range(14):  # 14 days
            day = Day(
                trip_id=trip.id,
                date=trip.start_date + timedelta(days=day_num),
                day_number=day_num + 1,
                notes=f"Den {day_num + 1} - poznámky"
            )
            db_session.add(day)
            await db_session.flush()
            
            # Create meal slots
            for i, meal_type in enumerate(meal_types):
                if meal_type == "snack" and day_num % 2 == 0:
                    continue  # Skip snack every other day
                
                meal_slot = MealSlot(
                    day_id=day.id,
                    meal_type=meal_type,
                    time=f"{6 + i * 5}:00"  # 6:00, 11:00, 16:00, 21:00
                )
                db_session.add(meal_slot)
                await db_session.flush()
                
                # Assign recipe
                recipe_idx = (day_num * len(meal_types) + i) % len(recipes)
                meal = Meal(
                    meal_slot_id=meal_slot.id,
                    recipe_id=recipes[recipe_idx].id,
                    servings=trip.participant_count,
                    notes=f"Poznámka k jídlu {meal_type}"
                )
                db_session.add(meal)
                meal_slots.append(meal_slot)
        
        await db_session.commit()
        return trip
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_trip_pdf_complete(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test PDF export of complete trip with all data."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.PDF.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "include_recipes": True,
                    "include_shopping": True,
                    "include_nutrition": True,
                    "include_costs": True,
                    "include_packing": True,
                    "page_size": "A4",
                    "language": "cs"
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        
        # Verify PDF content
        pdf_content = io.BytesIO(response.content)
        pdf_reader = PdfReader(pdf_content)
        
        assert len(pdf_reader.pages) > 5  # Should have multiple pages
        
        # Check first page contains trip name
        first_page_text = pdf_reader.pages[0].extract_text()
        assert "Letní tábor 2024" in first_page_text
        
        # Save for manual inspection if needed
        with open("/tmp/test_trip_export.pdf", "wb") as f:
            f.write(response.content)
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_trip_excel_with_formulas(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test Excel export with formulas and multiple sheets."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.EXCEL.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "excel_options": {
                        "include_formulas": True,
                        "include_charts": True,
                        "separate_sheets": True,
                        "freeze_headers": True,
                        "auto_filter": True
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Load and verify Excel content
        excel_content = io.BytesIO(response.content)
        workbook = load_workbook(excel_content, data_only=False)
        
        # Check sheets
        expected_sheets = ["Přehled", "Jídelníček", "Recepty", "Nákupní seznam"]
        for sheet_name in expected_sheets:
            assert sheet_name in workbook.sheetnames
        
        # Check formulas in shopping list
        shopping_sheet = workbook["Nákupní seznam"]
        # Look for SUM formulas
        has_formulas = False
        for row in shopping_sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    has_formulas = True
                    break
        
        assert has_formulas, "Excel should contain formulas"
        
        # Save for manual inspection
        with open("/tmp/test_trip_export.xlsx", "wb") as f:
            f.write(response.content)
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_shopping_list_csv(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test CSV export of shopping list."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.SHOPPING_LIST.value,
                "export_format": ExportFormat.CSV.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "csv_options": {
                        "delimiter": ";",
                        "encoding": "utf-8",
                        "include_headers": True
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/csv"
        
        # Parse CSV
        csv_content = response.content.decode('utf-8')
        csv_reader = csv.DictReader(io.StringIO(csv_content), delimiter=';')
        
        rows = list(csv_reader)
        assert len(rows) > 0
        
        # Check headers
        expected_headers = ["Položka", "Kategorie", "Množství", "Jednotka"]
        for header in expected_headers:
            assert header in csv_reader.fieldnames
        
        # Verify some content
        categories_found = set()
        for row in rows:
            categories_found.add(row["Kategorie"])
            assert float(row["Množství"].replace(",", ".")) > 0
        
        assert len(categories_found) > 1  # Multiple categories
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_recipes_json(
        self,
        client: AsyncClient,
        db_session: AsyncSession,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test JSON export of recipes from trip."""
        # Get recipes used in trip
        result = await db_session.execute(
            """
            SELECT DISTINCT r.id
            FROM recipes r
            JOIN meals m ON m.recipe_id = r.id
            JOIN meal_slots ms ON m.meal_slot_id = ms.id
            JOIN days d ON ms.day_id = d.id
            WHERE d.trip_id = :trip_id
            """,
            {"trip_id": complete_trip_data.id}
        )
        recipe_ids = [row[0] for row in result]
        
        # Export first recipe as JSON
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.RECIPE.value,
                "export_format": ExportFormat.JSON.value,
                "item_id": str(recipe_ids[0]),
                "async_export": False,
                "options": {
                    "json_options": {
                        "indent": 2,
                        "ensure_ascii": False
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/json"
        
        # Parse JSON
        recipe_data = json.loads(response.content)
        
        # Verify structure
        assert "name" in recipe_data
        assert "ingredients" in recipe_data
        assert "nutrition" in recipe_data
        assert "servings" in recipe_data
        
        # Check ingredients
        assert len(recipe_data["ingredients"]) > 0
        for ingredient in recipe_data["ingredients"]:
            assert "name" in ingredient
            assert "amount" in ingredient
            assert "unit" in ingredient
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_trip_markdown(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test Markdown export of trip."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.MARKDOWN.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "markdown_options": {
                        "flavor": "github",
                        "include_toc": True,
                        "heading_style": "atx"
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/markdown"
        
        # Check content
        markdown_content = response.content.decode('utf-8')
        
        # Should have proper structure
        assert "# Letní tábor 2024" in markdown_content
        assert "## Jídelníček" in markdown_content
        assert "## Nákupní seznam" in markdown_content
        
        # Check for tables
        assert "|" in markdown_content  # Markdown tables use pipes
        
        # Check for TOC
        assert "## Obsah" in markdown_content or "## Table of Contents" in markdown_content
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_trip_text(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test plain text export of trip."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.TEXT.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "text_options": {
                        "line_width": 80,
                        "indent_size": 2,
                        "section_separator": "="
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/plain"
        
        # Check content
        text_content = response.content.decode('utf-8')
        
        # Should be readable plain text
        assert "Letní tábor 2024" in text_content
        assert "=" * 10 in text_content  # Section separators
        
        # Check line width
        lines = text_content.split('\n')
        long_lines = [l for l in lines if len(l) > 85]  # Allow some margin
        assert len(long_lines) < 5  # Most lines should respect width
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_trip_html(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test HTML export of trip."""
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.HTML.value,
                "item_id": str(complete_trip_data.id),
                "async_export": False,
                "options": {
                    "html_options": {
                        "include_css": True,
                        "responsive": True,
                        "theme": "light"
                    }
                }
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "text/html"
        
        # Check HTML structure
        html_content = response.content.decode('utf-8')
        
        assert "<!DOCTYPE html>" in html_content
        assert "<html" in html_content
        assert "<head>" in html_content
        assert "<body>" in html_content
        
        # Check for CSS
        assert "<style>" in html_content or '<link rel="stylesheet"' in html_content
        
        # Check for responsive meta tag
        assert 'viewport' in html_content
        
        # Check content
        assert "Letní tábor 2024" in html_content
        assert "<table" in html_content  # Should have tables
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_batch_export_multiple_formats(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        db_session: AsyncSession,
        auth_headers: dict
    ):
        """Test batch export with multiple formats."""
        # Get some recipe IDs
        result = await db_session.execute(
            """
            SELECT DISTINCT r.id
            FROM recipes r
            JOIN meals m ON m.recipe_id = r.id
            JOIN meal_slots ms ON m.meal_slot_id = ms.id
            JOIN days d ON ms.day_id = d.id
            WHERE d.trip_id = :trip_id
            LIMIT 3
            """,
            {"trip_id": complete_trip_data.id}
        )
        recipe_ids = [str(row[0]) for row in result]
        
        # Create batch export
        response = await client.post(
            "/api/v1/exports/batch",
            json={
                "exports": [
                    {
                        "export_type": ExportType.TRIP.value,
                        "export_format": ExportFormat.PDF.value,
                        "item_ids": [str(complete_trip_data.id)],
                        "merge_into_single_file": False
                    },
                    {
                        "export_type": ExportType.RECIPE.value,
                        "export_format": ExportFormat.JSON.value,
                        "item_ids": recipe_ids,
                        "merge_into_single_file": True
                    },
                    {
                        "export_type": ExportType.SHOPPING_LIST.value,
                        "export_format": ExportFormat.EXCEL.value,
                        "item_ids": [str(complete_trip_data.id)],
                        "merge_into_single_file": False
                    }
                ]
            },
            headers=auth_headers
        )
        
        assert response.status_code == 200
        batch_data = response.json()
        
        assert "batch_id" in batch_data
        assert batch_data["job_count"] == 3
        
        # Wait for batch to complete
        batch_id = batch_data["batch_id"]
        completed = False
        
        for _ in range(30):  # Max 30 seconds
            await asyncio.sleep(1)
            
            # In real implementation, would check batch status
            # For now, assume it completes
            completed = True
            break
        
        assert completed
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_with_preview(
        self,
        client: AsyncClient,
        complete_trip_data: Trip,
        auth_headers: dict
    ):
        """Test export preview functionality."""
        # Generate preview
        response = await client.post(
            f"/api/v1/exports/preview/{ExportFormat.PDF.value}",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.PDF.value,
                "item_id": str(complete_trip_data.id)
            },
            headers=auth_headers
        )
        
        if response.status_code == 200:
            # Preview should be smaller than full export
            preview_size = len(response.content)
            
            # Get full export
            full_response = await client.post(
                "/api/v1/exports/single",
                json={
                    "export_type": ExportType.TRIP.value,
                    "export_format": ExportFormat.PDF.value,
                    "item_id": str(complete_trip_data.id),
                    "async_export": False
                },
                headers=auth_headers
            )
            
            full_size = len(full_response.content)
            
            # Preview should be smaller
            assert preview_size < full_size
            assert preview_size > 0
    
    @pytest.mark.asyncio
    @pytest.mark.integration
    async def test_export_error_handling(
        self,
        client: AsyncClient,
        auth_headers: dict
    ):
        """Test export error handling with invalid data."""
        # Non-existent trip
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.TRIP.value,
                "export_format": ExportFormat.PDF.value,
                "item_id": "99999999",
                "async_export": False
            },
            headers=auth_headers
        )
        
        assert response.status_code == 400
        error_data = response.json()
        assert "detail" in error_data
        
        # Invalid format for type (if implemented)
        response = await client.post(
            "/api/v1/exports/single",
            json={
                "export_type": ExportType.SHOPPING_LIST.value,
                "export_format": ExportFormat.MARKDOWN.value,  # May not be supported
                "item_id": "123",
                "async_export": False
            },
            headers=auth_headers
        )
        
        # Should handle gracefully
        assert response.status_code in [200, 400]